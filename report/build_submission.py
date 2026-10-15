"""Build the shareable submission/ folder: Markdown + DOCX report with identical content.

Usage:  python report/build_submission.py [--name "Hardik Kumar"] [--roll 2401MC39] [--zip] [--pdf]
Both files are rendered from ONE list of content blocks (`blocks()`), so they cannot drift apart.
Every number is read from results/ (metrics_aggregate.csv, selected_operating_point.json,
entity_timing.csv). Figures are copied from results/ (which evaluate_waitk.py plots from the CSVs).
--pdf converts the DOCX with Microsoft Word (Windows only) to check the page count.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import shutil
import sys
import zipfile
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from src.dataset import load_all  # noqa: E402
from src.entity_checker import category_of  # noqa: E402

RES, SUB = ROOT / "results", ROOT / "submission"
FIGS = ["latency_vs_entity_accuracy.png", "k_sweep.png", "pareto_frontier.png"]


def rcsv(p: Path) -> list[dict]:
    with open(p, encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def pct(x) -> str:
    return f"{100 * float(x):.1f}%"


def f2(x) -> str:
    return f"{float(x):.2f}"


# ------------------------------------------------------------------ content (single source)

def blocks(name: str, roll: str = "") -> list[tuple]:
    clips = load_all()
    agg = rcsv(RES / "metrics_aggregate.csv")
    by = {a["policy"]: a for a in agg}
    so = json.loads((RES / "selected_operating_point.json").read_text(encoding="utf-8"))
    p, sec = so["primary_words"], so["secondary_seconds_estimated"]
    et = rcsv(RES / "entity_timing.csv")
    k = p["selected"]["k"]
    sel = by[f"wait-{k}"]
    w1 = by["wait-1"]
    alts = p["alternatives"]
    n_ent = int(sel["n_entities"])
    n_src = sum(len(c.src_words) for c in clips)
    n_tgt = sum(len(c.tgt_tokens) for c in clips)
    cats = Counter(category_of(e) for c in clips for e in c.entities)
    measured = all(c.timing_source == "mms_fa_auto" for c in clips)
    est = "" if measured else " (estimated)"
    alpha_map = ", ".join(f"α={s['alpha']}→k={s['selected_k']}" for s in p["utility_sensitivity"])
    flips = "; ".join(f"k={a['selected_k']}→k={s['selected_k']} at α={s['alpha']}"
                      for a, s in zip(p["utility_sensitivity"], p["utility_sensitivity"][1:])
                      if a["selected_k"] != s["selected_k"])
    dom_w = {d["policy"]: d["dominates_waitk"] for d in so["reference_policies_dominating_waitk_words"]}
    dom_s = [d["policy"] for d in so["reference_policies_dominating_waitk_seconds_estimated"] if d["dominates_waitk"]]

    def cat_rate(pol, cat):
        rs = [e for e in et if e["policy"] == pol and e["category"] == cat]
        return f"{sum(e['safe'] == 'True' for e in rs)}/{len(rs)}"

    B: list[tuple] = []
    B.append(("title", "Where should a Hindi–English simultaneous translator wait? A text-based wait-k sweep"))
    who = f"{name} (Roll No. {roll})" if roll else name
    B.append(("meta", f"{who} · Inter IIT Tech Meet 15.0 · Sarvam.ai mini problem statement, ML track (Part E) · "
                      "code, tests and RESULTS.md in the project repository"))
    B.append(("note", "**Status of evidence.** This is a *text-based timing simulation* on six annotated Hindi text samples. "
                      "No recordings exist yet, no translation or speech model was trained or run, and every latency in seconds is an "
                      f"**estimate** from placeholder timings ({'measured' if measured else 'not measured'}). The 30-boundary audio "
                      "alignment audit is still pending. Transcripts are scripted (not yet verified against audio) and the English references "
                      "are LLM drafts pending human review."))

    B.append(("h1", "1. Objective and problem framing"))
    B.append(("p", "The Sarvam task streams code-mixed Hindi speech in and English speech out. Spoken English cannot be taken back, "
                   "and the system is scored on both entity accuracy (amounts, numbers, names, negation, corrections) and latency. "
                   "Hindi is verb-final, and speakers correct themselves (\"पाँच लाख… नहीं नहीं, पाँच करोड़\"), so speaking too early "
                   "commits to words whose deciding evidence has not been heard yet. Waiting longer costs latency. **Objective:** sweep "
                   "the wait-k policy over k = 1…10 on our annotated samples, measure the accuracy–latency trade-off with metrics that are "
                   "honest about what a text simulation can show, and choose an operating point with an explicit, reproducible rule."))

    B.append(("h1", "2. Relevant research: STACL and Hibiki"))
    B.append(("p", "**From the papers. STACL [2]** introduces prefix-to-prefix training and the *wait-k* policy: the model reads k "
                   "source words, then alternates one target word per source word, always k words behind. Its read schedule is "
                   "g(t) = min(k + t − 1, |x|) (Eq. 6), the number of source words read before target word t. STACL also defines "
                   "*Average Lagging* (Eq. 9): AL = (1/τ) Σ_{t≤τ} [g(t) − (t−1)/r], where r = |y|/|x| and τ is the first t at which g(t) = |x|. "
                   "AL is in source words, and AL = k when |x| = |y|. SimulEval [3] applies the same formula to time delays (seconds)."))
    B.append(("p", "**Hibiki [4]** (French→English speech-to-speech) does *not* use a fixed k. Its *contextual alignment* finds, for each "
                   "target word, the source word whose arrival most increases that word's log-likelihood under an off-the-shelf "
                   "translation model (MADLAD-3B, fed truncated source prefixes). It then builds training targets that lag this alignment "
                   "by at least δ_ctx = 2 s, inserting silence or using alignment-aware TTS. In its ablation, a constant 2 s lag reaches "
                   "ASR-BLEU 10.0 at LAAL 2.49 s, while contextual alignment reaches 26.6 at LAAL 5.0 s (sentence alignment: 25.6 at 21.49 s). "
                   "Our reading of that table: *where* the wait is placed matters, not only how long it is."))
    B.append(("p", "**Our implementation (what differs).** We implement and evaluate **wait-k only**, plus our own heuristic "
                   "entity-aware waiting (EAW) variants as a secondary comparison. We do **not** reproduce Hibiki. We run no "
                   "translation model, compute no likelihood differences and train nothing. Hibiki is the inspiration for one design choice: "
                   "each annotated entity records the Hindi word that determines it (`src_idx`), a hand-made, entity-only analogue of "
                   "\"the source word that makes the target word predictable\"."))

    B.append(("h1", "3. Dataset: the six existing Hindi text samples"))
    B.append(("p", f"Six scripted code-mixed Hindi passages (`data/annotations/clips.json`) contain {n_src} Hindi words and {n_tgt} "
                   f"English reference words. They have {n_ent} hand-annotated entities: " +
                   ", ".join(f"{v} {c}" for c, v in cats.most_common()) +
                   ". Each entity records (i) its determining Hindi word index, e.g. करोड़ in \"पाँच करोड़\" or the corrected digit; "
                   "(ii) the index of its first English word; and (iii) a hand-written *premature guess*, the best English available from "
                   "the prefix (\"five lakh\"), or `<unk>` for an unheard name. **Assumption:** source word timings are placeholder "
                   "estimates computed from character counts (clips of about 22–26 s). Each English word lasts a fixed 0.40 s. "
                   "Compute and TTS latency are 0 because nothing was run. Statuses remain `pending_recording`, `scripted_unverified` and "
                   "`llm_draft_unreviewed`."))

    B.append(("h1", "4. Method and evaluation setup"))
    B.append(("bullets", [
        "**Simulation.** The English text is the fixed fluent reference, and only its release times change with k. Every k uses "
        "the same clips, tokenisation, references, entity annotations, checker and conventions. A source word counts as read at its "
        "end time. The last English word waits for the full source (EOS), and English speech queues behind itself (Sarvam's playback "
        "rule: a word starts when decided or when the previous one ends, whichever is later).",
        "**(1) Fixed-reference entity accuracy:** the checker on the unmodified reference. It is identical for every k by construction.",
        "**(2) Evidence-safe rate:** the share of entities committed after their determining Hindi word was read. This is a *proxy* "
        "for whether enough evidence was available, not proof that a model would produce the right word.",
        "**(3) Simulated entity accuracy (stress test, not a model):** an entity committed early is replaced by its premature "
        "guess, and the checker re-scores the text. Every annotated guess fails the checker, so (3) equals (2) here "
        f"({'true' if so['sim_equals_evidence_safe'] else 'false'} for all {len(agg)} policies).",
        "**(4) AL in source words** (STACL Eq. 9), averaged over clips.",
        f"**(5) Latency in seconds{est}:** SimulEval AL of English speech onsets and End Offset (output end minus source end).",
        "**Selection rule** (implemented in `src/sweep.py`). The candidates are wait-1…wait-10. "
        "(a) Pareto frontier on (max evidence-safe rate, min AL in words): a point is dominated if another has at least the same rate "
        "and no greater AL, with one strict. (b) Knee: after min-max normalisation, the frontier point with the largest distance above "
        "the chord joining the frontier's extremes; ties go to the smaller k. The knee is flagged *weak* when that distance is smaller "
        "than one entity flipping (1/37 of the rate). (c) Secondary sensitivity: utility = α·rate_norm + (1−α)·(1 − AL_norm) "
        "for α = 0.1…0.9. α is a preference weight, not a constant.",
    ]))

    B.append(("h1", "5. Results: the k sweep"))
    rows = []
    for a in agg:
        if a["kind"] == "full":
            continue
        rows.append([a["policy"], f2(a["AL_words"]), f2(a["AL_onset_s"]), f2(a["end_offset_s"]),
                     pct(a["evidence_safe_rate"]), pct(a["sim_entity_acc"]), pct(a["fixed_text_acc"])])
    wc = by["whole-clip"]
    B.append(("table", ["Policy", "AL (words)", f"AL onset (s){est}", f"End Offset (s){est}",
                        "Evidence-safe", "Simulated acc.", "Fixed-ref. acc."], rows,
              f"Table 1. Means over 6 clips (latency); entity rates pooled over {n_ent} entities. EAW = our heuristic variants. "
              f"Whole-clip baseline (not shown): AL {f2(wc['AL_words'])} words, {pct(wc['evidence_safe_rate'])} evidence-safe. "
              "Source: results/metrics_aggregate.csv."))
    B.append(("p", f"**Observed.** Fixed-reference accuracy is **{pct(sel['fixed_text_acc'])} for every k**, so it cannot "
                   f"discriminate between policies. From wait-1 to wait-10, the evidence-safe rate rises from {pct(w1['evidence_safe_rate'])} "
                   f"to {pct(by['wait-10']['evidence_safe_rate'])}, while AL grows almost linearly, by about one source word per step "
                   f"({f2(w1['AL_words'])} → {f2(by['wait-10']['AL_words'])}). wait-10 adds latency with no safety gain over wait-9."))
    B.append(("fig", "latency_vs_entity_accuracy.png", 16.5,
              "Figure 1. Simulated entity accuracy vs AL in words (left, timing-independent) and vs the estimated AL of English "
              "speech onsets (right, placeholder timings). The dashed line is the fixed-reference accuracy, which is the same for every policy."))
    B.append(("fig", "k_sweep.png", 16.5,
              "Figure 2. Each metric against k. The entity rates and AL in words do not depend on timings; the seconds panel is estimated."))

    B.append(("h1", "6. Operating-point selection"))
    B.append(("fig", "pareto_frontier.png", 12.0,
              "Figure 3. Pareto frontier of the wait-k sweep (filled = non-dominated, hollow = dominated), the chord used for the knee "
              "test, the selected point and its alternatives. Generated from results/metrics_aggregate.csv."))
    B.append(("bullets", [
        f"**Pareto frontier:** k = {', '.join(map(str, p['pareto_k']))}; dominated: k = {', '.join(map(str, p['dominated_k'])) or 'none'}.",
        f"**Selected (knee): wait-{k}**, with an evidence-safe rate of {pct(sel['evidence_safe_rate'])}, AL {f2(sel['AL_words'])} source words, "
        f"and an estimated AL onset of {f2(sel['AL_onset_s'])} s. Compared with wait-1, AL rises by {p['vs_most_aggressive']['latency_change']:.2f} "
        f"words and evidence-safe by {p['vs_most_aggressive']['accuracy_change_points']:.1f} points.",
        "**Near-equivalent alternatives:** " + "; ".join(
            f"wait-{a['k']} ({pct(a['evidence_safe_rate'])}, AL {f2(a['AL_words'])})" for a in alts) + ".",
        f"**The knee is {'weak' if p['knee_is_weak'] else 'clear'}:** its distance above the chord ({p['knee_distance']}) is "
        f"{'smaller' if p['knee_is_weak'] else 'larger'} than the effect of one entity flipping ({p['one_entity_in_normalised_units']}). "
        "The curve is close to linear, so the data do not single out one k.",
        f"**Sensitivity:** {alpha_map}. The selected k switches abruptly ({flips}) instead of drifting "
        "gradually, which is the signature of a near-linear trade-off. On the estimated seconds axis the knee is again "
        f"wait-{sec['selected']['k']}.",
        "**Reference policies:** on AL in words, " + "; ".join(
            f"{pol} dominates wait-{', '.join(map(str, v))}" for pol, v in dom_w.items() if v) +
        f". On estimated seconds, {'none dominate' if not dom_s else ', '.join(dom_s) + ' dominate'}, because EAW-pause pays "
        "for observed pauses and sentence-level output queues long English speech.",
    ]))
    B.append(("p", f"**Interpretation.** We recommend **wait-{k} as the nominal operating point, with wait-{alts[0]['k']}–wait-{alts[-1]['k']} "
                   "as an equivalent band** and a starting point for tuning on real audio. This is a preference on a near-linear curve, "
                   "not an optimum. Safety per category explains why fixed k is blunt. Under wait-3, corrections are "
                   f"{cat_rate('wait-3', 'correction')} safe and negations {cat_rate('wait-3', 'negation')}; under wait-{k}, they are "
                   f"{cat_rate(f'wait-{k}', 'correction')} and {cat_rate(f'wait-{k}', 'negation')}. Negation needs the clause-final Hindi verb, "
                   "and no small fixed k reliably reaches it. This is the content-dependent waiting that Hibiki's contextual alignment learns "
                   "from data."))

    B.append(("h1", "7. Limitations and next steps"))
    B.append(("bullets", [
        "**Not a model measurement.** The reference translation is fixed, so a prefix-to-prefix model's real errors, both "
        "hallucinated anticipation and good anticipation, are not captured. Rates (2) and (3) are proxies.",
        f"**Small sample.** Six scripted text samples, one author, {n_ent} entities: one entity is {100 / n_ent:.1f} points, so "
        "adjacent k values are not statistically distinguishable. The premature guesses and `src_idx` labels are the author's own.",
        "**Seconds are estimates.** Real word durations, pauses, English TTS durations, compute latency and playback queues will "
        "change the seconds axis. Longer real TTS words would raise the onset AL and End Offset for every k, and most of all for "
        "large k and sentence-level output. Word-based AL and the evidence-safe rate will change only if transcripts change.",
        "**Pending requirements:** six natural recordings (20–40 s), verbatim transcripts verified against audio, forced alignment, "
        "the 30-boundary manual audit, and human review of the translations and checker labels.",
        "**Future work:** re-run the sweep on aligned recordings; replace fixed k with a target-triggered or contextual-alignment "
        "(Hibiki-style) policy using a small MT model's likelihood gains; evaluate a real streaming model's output with the same checker.",
    ]))

    B.append(("h1", "References"))
    B.append(("refs", [
        "[1] Sarvam AI. Inter IIT Tech Meet 15.0 problem statement and Sarvam.ai mini problem statement (provided PDFs).",
        "[2] M. Ma et al. STACL: Simultaneous Translation with Implicit Anticipation and Controllable Latency using "
        "Prefix-to-Prefix Framework. ACL 2019, pp. 3025–3036 (Eq. 6 wait-k, Eq. 9 Average Lagging).",
        "[3] X. Ma, M. J. Dousti, C. Wang, J. Gu, J. Pino. SimulEval: An Evaluation Toolkit for Simultaneous Translation. "
        "EMNLP 2020 System Demonstrations, pp. 144–150; ALScorer in simuleval/evaluator/scorers/latency_scorer.py.",
        "[4] T. Labiausse, L. Mazaré, E. Grave, A. Défossez, N. Zeghidour. High-Fidelity Simultaneous Speech-To-Speech "
        "Translation (Hibiki). ICML 2025, PMLR 267 (Sec. 3.2 contextual alignment; Table 4 ablation).",
        "[5] S. Papi, M. Gaido, M. Negri, M. Turchi. Over-Generation Cannot Be Rewarded: Length-Adaptive Average Lagging (LAAL). "
        "AutoSimTrans 2022, pp. 12–17.",
    ]))
    return B


# ------------------------------------------------------------------ renderers

def to_markdown(B: list[tuple]) -> str:
    L = []
    for b in B:
        kind = b[0]
        if kind == "title":
            L += [f"# {b[1]}", ""]
        elif kind == "meta":
            L += [f"*{b[1]}*", ""]
        elif kind == "note":
            L += [f"> {b[1]}", ""]
        elif kind == "h1":
            L += [f"## {b[1]}", ""]
        elif kind == "p":
            L += [b[1], ""]
        elif kind == "bullets":
            L += [f"- {x}" for x in b[1]] + [""]
        elif kind == "table":
            _, head, rows, cap = b
            L += ["| " + " | ".join(head) + " |", "|" + "---|" * len(head)]
            L += ["| " + " | ".join(r) + " |" for r in rows]
            L += ["", f"*{cap}*", ""]
        elif kind == "fig":
            _, fn, _, cap = b
            L += [f"![{cap.split('.')[0]}](figures/{fn})", "", f"*{cap}*", ""]
        elif kind == "refs":
            L += [f"{r}  " for r in b[1]] + [""]
    return "\n".join(L)


INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def add_runs(par, text: str, size=None, italic=False):
    from docx.shared import Pt
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith("*"):
            r = par.add_run(part[1:-1]); r.italic = True
        elif part.startswith("`"):
            r = par.add_run(part[1:-1]); r.font.name = "Consolas"
        else:
            r = par.add_run(part)
        if italic:
            r.italic = True
        if size:
            r.font.size = Pt(size)


AUTHOR = "Hardik Kumar"


def to_docx(B: list[tuple], out: Path) -> None:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(1.8))
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(9.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Nirmala UI")
    st.element.rPr.rFonts.set(qn("w:cs"), "Nirmala UI")
    st.paragraph_format.space_after = Pt(3)
    st.paragraph_format.line_spacing = 1.0
    for h, size in (("Heading 1", 11.5), ("Heading 2", 10)):
        hs = doc.styles[h]
        hs.font.size, hs.font.name = Pt(size), "Calibri"
        hs.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        hs.paragraph_format.space_before, hs.paragraph_format.space_after = Pt(6), Pt(2)
    for b in B:
        kind = b[0]
        if kind == "title":
            p = doc.add_heading(level=0)
            add_runs(p, b[1]); [setattr(r.font, "size", Pt(15)) for r in p.runs]
        elif kind == "meta":
            add_runs(doc.add_paragraph(), b[1], size=8.5, italic=True)
        elif kind == "note":
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "FDF2E9")
            pPr.append(shd)
            add_runs(p, b[1], size=8.5)
        elif kind == "h1":
            doc.add_heading(b[1], level=1)
        elif kind == "p":
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_runs(p, b[1])
        elif kind == "bullets":
            for x in b[1]:
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(1)
                add_runs(p, x)
        elif kind == "table":
            _, head, rows, cap = b
            t = doc.add_table(rows=1 + len(rows), cols=len(head))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, row in enumerate([head] + rows):
                for j, val in enumerate(row):
                    cell = t.cell(i, j)
                    cell.text = ""
                    par = cell.paragraphs[0]
                    par.paragraph_format.space_after = Pt(0)
                    r = par.add_run(val); r.font.size = Pt(7.5); r.bold = i == 0
                    if j:
                        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    if i == 0:
                        tcPr = cell._tc.get_or_add_tcPr()
                        shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "E8EEF5")
                        tcPr.append(shd)
            add_runs(doc.add_paragraph(), cap, size=8, italic=True)
        elif kind == "fig":
            _, fn, width_cm, cap = b
            doc.add_picture(str(SUB / "figures" / fn), width=Cm(width_cm))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(doc.add_paragraph(), cap, size=8, italic=True)
        elif kind == "refs":
            for r in b[1]:
                p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(1)
                add_runs(p, r, size=8)
    doc.core_properties.title = "Text-based wait-k sweep (Sarvam mini PS, ML track)"
    doc.core_properties.author = AUTHOR
    doc.save(out)


def docx_to_pdf(docx: Path, pdf: Path) -> int:
    """Render with Microsoft Word (COM) and return the page count."""
    import subprocess
    ps = (f"$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
          f"$d = $w.Documents.Open('{docx}'); $d.SaveAs([ref]'{pdf}', [ref]17); $d.Close(); $w.Quit()")
    subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True)
    import pypdfium2 as pdfium
    return len(pdfium.PdfDocument(str(pdf)))


def write_readme(name: str, roll: str, pages: int | None) -> None:
    so = json.loads((RES / "selected_operating_point.json").read_text(encoding="utf-8"))["primary_words"]
    (SUB / "README.md").write_text(f"""# Submission package: text-based wait-k sweep (Sarvam mini PS, ML track)

Generated by `.venv\\Scripts\\python report/build_submission.py` from the files in `results/`.
Do not edit the numbers by hand. Re-run the commands below instead. The prose in `submission_report.docx`
and `submission_report.md` can be edited freely; both contain the same content.

## Files to upload

| File | What it contains |
|---|---|
| `submission_report.docx` | Editable Word report (figures embedded){f'. Renders to {pages} A4 pages in Word' if pages else ''} |
| `submission_report.md` | The same report in Markdown (figures linked from `figures/`) |
| `submission_report.pdf` | PDF rendered from the DOCX by Microsoft Word (page-count check; regenerate after editing the DOCX) |
| `RESULTS.md` | Every reported number with the command that produced it (copy of the repository's RESULTS.md) |
| `figures/latency_vs_entity_accuracy.png` | Main trade-off graph (simulated entity accuracy vs AL in words / estimated seconds) |
| `figures/k_sweep.png` | Entity metrics and latency as functions of k = 1–10 |
| `figures/pareto_frontier.png` | Pareto frontier, selected point (wait-{so['selected']['k']}) and alternatives |
| `data/metrics_aggregate.csv` | One row per policy (compact) |
| `data/metrics.csv` | One row per (clip, policy) |
| `data/k_sensitivity.csv` | Selected k under utility weights alpha = 0.1–0.9 |
| `data/selected_operating_point.json` | Machine-readable selection record and rationale |

`submission_package.zip` (one level up, in the repository root) contains exactly this folder.

## What these results are and are not

- A **text-based timing simulation** on six annotated Hindi text samples. No recordings and no trained model are involved.
- Latencies in seconds are **estimates** (placeholder source timings, 0.40 s per English word, compute = 0).
- Fixed-reference entity accuracy is 100% for every k by construction. The evidence-safe rate and the simulated accuracy are proxies.
- Pending: six natural recordings, transcript verification, forced alignment, the 30-boundary audit, and human review of translations.

## Reproduce (from the repository root)

```powershell
.venv\\Scripts\\python -m pytest -q
.venv\\Scripts\\python scripts/evaluate_waitk.py
.venv\\Scripts\\python scripts/write_results_md.py
.venv\\Scripts\\python report/build_submission.py --name "{name}" --roll "{roll}" --pdf --zip
```
""", encoding="utf-8")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--name", default="Hardik Kumar")
    ap.add_argument("--roll", default="2401MC39")
    ap.add_argument("--zip", action="store_true")
    ap.add_argument("--pdf", action="store_true", help="render the DOCX to PDF with MS Word and count pages")
    a = ap.parse_args()
    (SUB / "figures").mkdir(parents=True, exist_ok=True)
    (SUB / "data").mkdir(parents=True, exist_ok=True)
    for fn in FIGS:
        shutil.copy2(RES / fn, SUB / "figures" / fn)
    for fn in ("metrics_aggregate.csv", "metrics.csv", "k_sensitivity.csv", "selected_operating_point.json"):
        shutil.copy2(RES / fn, SUB / "data" / fn)
    shutil.copy2(ROOT / "RESULTS.md", SUB / "RESULTS.md")
    global AUTHOR
    AUTHOR = a.name
    B = blocks(a.name, a.roll)
    (SUB / "submission_report.md").write_text(to_markdown(B), encoding="utf-8")
    to_docx(B, SUB / "submission_report.docx")
    pages = None
    if a.pdf:
        pages = docx_to_pdf(SUB / "submission_report.docx", SUB / "submission_report.pdf")
        print(f"DOCX renders to {pages} pages in Word" + ("  ** OVER 4 PAGES **" if pages > 4 else ""))
    write_readme(a.name, a.roll, pages)
    if a.zip:
        z = ROOT / "submission_package.zip"
        with zipfile.ZipFile(z, "w", zipfile.ZIP_DEFLATED) as zf:
            for f in sorted(SUB.rglob("*")):
                if f.is_file() and "__pycache__" not in f.parts and not f.name.startswith("~$"):
                    zf.write(f, Path("submission") / f.relative_to(SUB))
        print(f"wrote {z} ({z.stat().st_size // 1024} KB)")
    print(f"wrote {SUB}")
    return 0 if not pages or pages <= 4 else 1


if __name__ == "__main__":
    sys.exit(main())
